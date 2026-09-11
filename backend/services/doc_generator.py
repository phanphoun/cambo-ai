"""Document generation engine for creating publication-grade PDF, Word (.docx), Markdown, and CSV documents with full Khmer Unicode & LaTeX Math support."""
import os
import io
import re
import time
import uuid
import json
import shutil
import logging
import tempfile
import subprocess
from pathlib import Path
from typing import Optional, Dict, Any, List

from markdown_it import MarkdownIt

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

logger = logging.getLogger("cambo.services.doc_generator")

GENERATED_DOCS_DIR = Path(__file__).resolve().parent.parent / "data" / "generated_docs"
GENERATED_DOCS_DIR.mkdir(parents=True, exist_ok=True)
DOCS_META_FILE = GENERATED_DOCS_DIR / "metadata.json"

FONTS_DIR = Path(__file__).resolve().parent.parent / "data" / "fonts"
FONTS_DIR.mkdir(parents=True, exist_ok=True)
KHMER_FONT_PATH = FONTS_DIR / "NotoSansKhmer.ttf"

KATEX_DIR = Path(__file__).resolve().parent.parent / "data" / "katex"

# Detect Google Chrome or Chromium binary for publication-grade headless PDF rendering
CHROME_BIN = (
    shutil.which("google-chrome")
    or shutil.which("chromium")
    or shutil.which("chromium-browser")
    or os.environ.get("CHROME_PATH")
)

# Register Unicode Font for ReportLab fallback
REPORTLAB_FONT_NAME = "Helvetica"
if KHMER_FONT_PATH.exists():
    try:
        pdfmetrics.registerFont(TTFont("NotoSansKhmer", str(KHMER_FONT_PATH)))
        REPORTLAB_FONT_NAME = "NotoSansKhmer"
        logger.info("Registered NotoSansKhmer font for ReportLab fallback.")
    except Exception as e:
        logger.error("Failed to register NotoSansKhmer font: %s", e)


def _clean_math_for_text(text: str) -> str:
    """Converts common LaTeX formulas into clean, readable Unicode math symbols for Word & plain text."""
    # Remove math delimiters
    cleaned = re.sub(r"\$\$(.*?)\$\$", r"\1", text)
    cleaned = re.sub(r"\$(.*?)\$", r"\1", cleaned)
    cleaned = re.sub(r"\\\[(.*?)\\\]", r"\1", cleaned)
    cleaned = re.sub(r"\\\((.*?)\\\)", r"\1", cleaned)

    replacements = [
        (r"\\vec\{([^\}]+)\}", r"\1⃗"),
        (r"\\sqrt\{([^\}]+)\}", r"√(\1)"),
        (r"\\frac\{([^\}]+)\}\{([^\}]+)\}", r"(\1 / \2)"),
        (r"\\left\(", "("),
        (r"\\right\)", ")"),
        (r"\\left\[", "["),
        (r"\\right\]", "]"),
        (r"\\div", "÷"),
        (r"\\times", "×"),
        (r"\\cdot", "·"),
        (r"\\pm", "±"),
        (r"\\approx", "≈"),
        (r"\\neq", "≠"),
        (r"\\le(q)?\b", "≤"),
        (r"\\ge(q)?\b", "≥"),
        (r"\\infty", "∞"),
        (r"\\sum", "∑"),
        (r"\\int", "∫"),
        (r"\\alpha", "α"),
        (r"\\beta", "β"),
        (r"\\gamma", "γ"),
        (r"\\pi", "π"),
        (r"\\theta", "θ"),
        (r"\\lambda", "λ"),
        (r"\\in\b", "∈"),
        (r"\\notin\b", "∉"),
        (r"\\,", " "),
        (r"\^2", "²"),
        (r"\^3", "³"),
        (r"\\begin\{[a-zA-Z*]+\}", ""),
        (r"\\end\{[a-zA-Z*]+\}", ""),
    ]
    for pattern, repl in replacements:
        cleaned = re.sub(pattern, repl, cleaned)

    return cleaned.strip()


class NumberedCanvas(canvas.Canvas):
    """Custom canvas that adds page numbers and Sastra AI footer branding for ReportLab fallback."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count: int):
        self.saveState()
        self.setFont(REPORTLAB_FONT_NAME, 8)
        self.setFillColor(colors.HexColor("#78716C"))
        
        self.setStrokeColor(colors.HexColor("#D6D3D1"))
        self.setLineWidth(0.5)
        self.line(54, 40, letter[0] - 54, 40)
        
        footer_text = f"Sastra AI Sovereign Assistant  •  Page {self._pageNumber} of {page_count}"
        self.drawString(54, 26, footer_text)
        
        date_str = time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime())
        self.drawRightString(letter[0] - 54, 26, date_str)
        self.restoreState()


class DocumentGenerator:
    """Orchestrates publication-grade generation of user-requested documents (PDF, Word, Markdown, CSV)."""

    def __init__(self):
        self._meta: Dict[str, dict] = {}
        self._load_meta()
        self._md_parser = MarkdownIt("commonmark").enable(["table", "strikethrough"])

    def _load_meta(self):
        if DOCS_META_FILE.exists():
            try:
                with open(DOCS_META_FILE, "r", encoding="utf-8") as f:
                    self._meta = json.load(f)
            except Exception as e:
                logger.error("Failed to load doc metadata: %s", e)

    def _save_meta(self):
        try:
            with open(DOCS_META_FILE, "w", encoding="utf-8") as f:
                json.dump(self._meta, f, indent=2)
        except Exception as e:
            logger.error("Failed to save doc metadata: %s", e)

    @staticmethod
    def _clean_or_extract_title(title: str, content: str) -> str:
        """Derives a meaningful, professional document title instead of generic fallback names."""
        clean_title = (title or "").strip()
        if clean_title and clean_title not in ("Sastra_AI_Response", "document", "Untitled", "Export", "Response"):
            return clean_title.replace("_", " ").strip()

        # Try to find primary markdown heading # Title or ## Title
        heading_match = re.search(r"^#+\s+([^\n]+)", content, re.MULTILINE)
        if heading_match:
            candidate = re.sub(r"[*_`#]", "", heading_match.group(1)).strip()
            if candidate and len(candidate) > 2:
                return candidate[:75]

        # Try to find first bold title **Title**
        bold_match = re.search(r"^\s*\*\*([^\*\n]+)\*\*", content, re.MULTILINE)
        if bold_match:
            candidate = bold_match.group(1).strip()
            if candidate and len(candidate) > 2:
                return candidate[:75]

        return "Sastra AI Document"

    def generate_pdf(
        self,
        title: str,
        content: str,
        author: str = "Sastra AI Assistant",
        subtitle: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Generates a publication-grade multi-page PDF document with full Khmer OpenType shaping & KaTeX math."""
        doc_id = f"doc_{uuid.uuid4().hex[:10]}"
        resolved_title = self._clean_or_extract_title(title, content)
        filename = f"{self._sanitize_filename(resolved_title)}.pdf"
        file_path = GENERATED_DOCS_DIR / f"{doc_id}.pdf"

        # 1. Primary Engine: Headless Chrome with full HarfBuzz Khmer shaping & KaTeX
        if CHROME_BIN:
            try:
                self._generate_pdf_chrome(
                    file_path=file_path,
                    title=resolved_title,
                    content=content,
                    author=author,
                    subtitle=subtitle or "Official Knowledge & Research Document",
                )
                if file_path.exists() and file_path.stat().st_size > 1024:
                    file_size = file_path.stat().st_size
                    meta = {
                        "id": doc_id,
                        "title": resolved_title,
                        "format": "pdf",
                        "filename": filename,
                        "size_bytes": file_size,
                        "size_formatted": self._format_bytes(file_size),
                        "created_at": time.time(),
                        "download_url": f"/api/documents/generated/{doc_id}",
                    }
                    self._meta[doc_id] = meta
                    self._save_meta()
                    return meta
            except Exception as e:
                logger.error("Headless Chrome PDF generation failed, falling back to ReportLab: %s", e)

        # 2. Fallback Engine: ReportLab with clean math & dividers
        self._generate_pdf_reportlab(
            file_path=file_path,
            title=resolved_title,
            content=content,
            author=author,
            subtitle=subtitle,
        )

        file_size = file_path.stat().st_size
        meta = {
            "id": doc_id,
            "title": resolved_title,
            "format": "pdf",
            "filename": filename,
            "size_bytes": file_size,
            "size_formatted": self._format_bytes(file_size),
            "created_at": time.time(),
            "download_url": f"/api/documents/generated/{doc_id}",
        }
        self._meta[doc_id] = meta
        self._save_meta()
        return meta

    def _generate_pdf_chrome(
        self,
        file_path: Path,
        title: str,
        content: str,
        author: str,
        subtitle: str,
    ):
        """Builds an HTML document with KaTeX and renders it to PDF via headless Google Chrome."""
        body_html = self._md_parser.render(content.strip())

        # Determine KaTeX stylesheet and script locations
        katex_css_uri = "https://cdn.jsdelivr.net/npm/katex@0.16.21/dist/katex.min.css"
        katex_js_uri = "https://cdn.jsdelivr.net/npm/katex@0.16.21/dist/katex.min.js"
        auto_render_js_uri = "https://cdn.jsdelivr.net/npm/katex@0.16.21/dist/contrib/auto-render.min.js"

        local_css = KATEX_DIR / "katex.min.css"
        local_js = KATEX_DIR / "katex.min.js"
        local_auto = KATEX_DIR / "contrib" / "auto-render.min.js"

        if local_css.exists() and local_js.exists() and local_auto.exists():
            katex_css_uri = local_css.as_uri()
            katex_js_uri = local_js.as_uri()
            auto_render_js_uri = local_auto.as_uri()

        date_str = time.strftime("%B %d, %Y")

        html_template = f"""<!DOCTYPE html>
<html lang="km">
<head>
<meta charset="utf-8">
<title>{title}</title>
<link rel="stylesheet" href="{katex_css_uri}">
<script src="{katex_js_uri}"></script>
<script src="{auto_render_js_uri}"></script>
<style>
  @page {{
    size: A4;
    margin: 22mm 18mm 22mm 18mm;
    @bottom-left {{
      content: "Sastra AI Sovereign Assistant";
      font-size: 8pt;
      color: #78716c;
      font-family: "Kantumruy Pro", "Noto Sans Khmer", "Inter", sans-serif;
    }}
    @bottom-right {{
      content: "Page " counter(page) " of " counter(pages);
      font-size: 8pt;
      color: #78716c;
      font-family: "Inter", sans-serif;
    }}
  }}
  * {{
    box-sizing: border-box;
  }}
  body {{
    font-family: "Kantumruy Pro", "Noto Sans Khmer", "Inter", sans-serif;
    font-size: 10.5pt;
    line-height: 1.85;
    color: #1c1917;
    background: #ffffff;
    margin: 0;
    padding: 0;
    text-rendering: optimizeLegibility;
  }}
  /* Header branding bar */
  .doc-header {{
    padding-bottom: 12px;
    margin-bottom: 22px;
    border-bottom: 2px solid #b45309;
    display: flex;
    justify-content: space-between;
    align-items: flex-end;
  }}
  .brand-title {{
    font-size: 16pt;
    font-weight: 700;
    color: #92400e;
    margin: 0;
    line-height: 1.25;
  }}
  .brand-subtitle {{
    font-size: 8.5pt;
    color: #78716c;
    margin-top: 4px;
    letter-spacing: 0.3px;
  }}
  .doc-date {{
    font-size: 8pt;
    color: #a8a29e;
    text-align: right;
  }}

  /* Typography */
  h1 {{
    font-size: 16pt;
    font-weight: 700;
    color: #92400e;
    margin-top: 22px;
    margin-bottom: 8px;
    line-height: 1.35;
    page-break-after: avoid;
  }}
  h2 {{
    font-size: 13pt;
    font-weight: 600;
    color: #b45309;
    margin-top: 18px;
    margin-bottom: 6px;
    line-height: 1.4;
    page-break-after: avoid;
  }}
  h3 {{
    font-size: 11.5pt;
    font-weight: 600;
    color: #1c1917;
    margin-top: 14px;
    margin-bottom: 4px;
    page-break-after: avoid;
  }}
  p {{
    margin-top: 6px;
    margin-bottom: 10px;
    text-align: justify;
  }}
  ul, ol {{
    margin-top: 4px;
    margin-bottom: 10px;
    padding-left: 20px;
  }}
  li {{
    margin-bottom: 4px;
  }}
  hr {{
    border: none;
    border-top: 1px solid #e7e5e4;
    margin: 18px 0;
  }}
  blockquote {{
    margin: 12px 0;
    padding: 10px 16px;
    border-left: 3.5px solid #d97706;
    background: #fffbeb;
    border-radius: 0 8px 8px 0;
    color: #78350f;
    page-break-inside: avoid;
  }}
  strong {{
    font-weight: 700;
    color: #0f172a;
  }}
  pre {{
    background: #f5f5f4;
    border: 1px solid #e7e5e4;
    border-radius: 8px;
    padding: 12px;
    overflow-x: auto;
    font-size: 9pt;
    line-height: 1.5;
    page-break-inside: avoid;
  }}
  code {{
    font-family: "JetBrains Mono", "Courier New", monospace;
    font-size: 0.9em;
    background: #fef3c7;
    color: #92400e;
    padding: 2px 5px;
    border-radius: 4px;
    border: 1px solid #fde68a;
  }}
  pre code {{
    background: transparent;
    border: none;
    padding: 0;
    color: #1c1917;
  }}

  /* KaTeX Math block styling */
  .katex-display {{
    margin: 14px 0 !important;
    padding: 12px 16px;
    background: #fdfaf4;
    border: 1px solid #f3e8d2;
    border-radius: 8px;
    text-align: center;
    page-break-inside: avoid;
  }}
  .katex {{
    font-size: 1.05em;
  }}

  /* Tables */
  table {{
    width: 100%;
    border-collapse: collapse;
    margin: 16px 0;
    font-size: 9.5pt;
    page-break-inside: avoid;
  }}
  th {{
    background: #92400e;
    color: #ffffff;
    font-weight: 600;
    padding: 8px 12px;
    text-align: left;
    border: 1px solid #78350f;
  }}
  td {{
    padding: 7px 12px;
    border: 1px solid #e7e5e4;
  }}
  tr:nth-child(even) td {{
    background: #fafaf9;
  }}
</style>
</head>
<body>
  <div class="doc-header">
    <div>
      <div class="brand-title">Sastra AI Sovereign Document</div>
      <div class="brand-subtitle">{subtitle}  •  Authored by {author}</div>
    </div>
    <div class="doc-date">{date_str}</div>
  </div>

  <div class="doc-content">
    {f"<h1>{title}</h1>" if "<h1" not in body_html else ""}
    {body_html}
  </div>

  <script>
    renderMathInElement(document.body, {{
      delimiters: [
        {{left: "$$", right: "$$", display: true}},
        {{left: "\\\\[", right: "\\\\]", display: true}},
        {{left: "$", right: "$", display: false}},
        {{left: "\\\\(", right: "\\\\)", display: false}}
      ],
      throwOnError: false
    }});
  </script>
</body>
</html>
"""
        with tempfile.NamedTemporaryFile("w", suffix=".html", encoding="utf-8", delete=False) as tmp_file:
            tmp_file.write(html_template)
            tmp_html_path = tmp_file.name

        try:
            cmd = [
                CHROME_BIN,
                "--headless=new",
                "--disable-gpu",
                "--no-sandbox",
                "--no-pdf-header-footer",
                "--allow-file-access-from-files",
                "--run-all-compositor-stages-before-draw",
                "--virtual-time-budget=2500",
                f"--print-to-pdf={file_path}",
                tmp_html_path,
            ]
            subprocess.run(cmd, check=True, timeout=30, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        finally:
            try:
                os.remove(tmp_html_path)
            except OSError:
                pass

    def _generate_pdf_reportlab(
        self,
        file_path: Path,
        title: str,
        content: str,
        author: str,
        subtitle: Optional[str] = None,
    ):
        """Fallback PDF generator using ReportLab with clean math and divider support."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54,
        )

        styles = getSampleStyleSheet()
        gold_color = colors.HexColor("#996B00")
        dark_color = colors.HexColor("#1C1917")

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontName=REPORTLAB_FONT_NAME,
            fontSize=18,
            leading=24,
            textColor=dark_color,
            spaceAfter=4,
        )
        subtitle_style = ParagraphStyle(
            "DocSubtitle",
            parent=styles["Normal"],
            fontName=REPORTLAB_FONT_NAME,
            fontSize=9.5,
            leading=14,
            textColor=gold_color,
            spaceAfter=10,
        )
        h1_style = ParagraphStyle(
            "DocH1",
            parent=styles["Heading1"],
            fontName=REPORTLAB_FONT_NAME,
            fontSize=14,
            leading=19,
            textColor=gold_color,
            spaceBefore=14,
            spaceAfter=6,
        )
        h2_style = ParagraphStyle(
            "DocH2",
            parent=styles["Heading2"],
            fontName=REPORTLAB_FONT_NAME,
            fontSize=12,
            leading=16,
            textColor=dark_color,
            spaceBefore=10,
            spaceAfter=4,
        )
        body_style = ParagraphStyle(
            "DocBody",
            parent=styles["Normal"],
            fontName=REPORTLAB_FONT_NAME,
            fontSize=9.5,
            leading=15,
            textColor=dark_color,
            spaceAfter=6,
        )
        bullet_style = ParagraphStyle(
            "DocBullet",
            parent=styles["Normal"],
            fontName=REPORTLAB_FONT_NAME,
            fontSize=9.5,
            leading=14,
            textColor=dark_color,
            leftIndent=14,
            spaceAfter=3,
        )
        cell_style = ParagraphStyle(
            "DocCell",
            parent=styles["Normal"],
            fontName=REPORTLAB_FONT_NAME,
            fontSize=8.5,
            leading=12,
            textColor=dark_color,
        )
        cell_header_style = ParagraphStyle(
            "DocCellHeader",
            parent=styles["Normal"],
            fontName=REPORTLAB_FONT_NAME,
            fontSize=8.5,
            leading=12,
            textColor=colors.white,
        )

        story = []
        story.append(Paragraph(self._format_markdown_for_reportlab(title), title_style))
        if subtitle:
            story.append(Paragraph(self._format_markdown_for_reportlab(subtitle), subtitle_style))
        else:
            story.append(Paragraph(f"Authored by {author}  •  {time.strftime('%B %d, %Y')}", subtitle_style))

        story.append(HRFlowable(width="100%", thickness=1.5, color=gold_color, spaceBefore=2, spaceAfter=12))

        lines = content.strip().split("\n")
        table_buffer: List[str] = []

        def flush_table():
            if not table_buffer:
                return
            table_data = []
            for row_idx, t_line in enumerate(table_buffer):
                if re.match(r"^\|?\s*[-:]+\s*\|", t_line):
                    continue
                cells = [c.strip() for c in t_line.split("|")[1:-1]]
                if not cells:
                    continue
                row_flowables = []
                for cell_text in cells:
                    cell_xml = self._format_markdown_for_reportlab(_clean_math_for_text(cell_text))
                    style_to_use = cell_header_style if row_idx == 0 else cell_style
                    row_flowables.append(Paragraph(cell_xml, style_to_use))
                table_data.append(row_flowables)

            if table_data:
                col_count = len(table_data[0])
                avail_width = letter[0] - 108
                col_width = avail_width / max(col_count, 1)
                t = Table(table_data, colWidths=[col_width] * col_count)
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#996B00")),
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F8F6")]),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E7E5E4")),
                ]))
                story.append(Spacer(1, 4))
                story.append(t)
                story.append(Spacer(1, 8))
            table_buffer.clear()

        for line in lines:
            line_str = line.strip()
            if line_str.startswith("|") and line_str.endswith("|"):
                table_buffer.append(line_str)
                continue
            else:
                flush_table()

            if not line_str:
                story.append(Spacer(1, 4))
            elif re.match(r"^[-*_]{3,}$", line_str):
                # Horizontal rule separator
                story.append(HRFlowable(width="100%", thickness=0.75, color=colors.HexColor("#E7E5E4"), spaceBefore=8, spaceAfter=8))
            elif line_str.startswith("### "):
                clean_heading = self._format_markdown_for_reportlab(_clean_math_for_text(line_str[4:]))
                story.append(Paragraph(clean_heading, h2_style))
            elif line_str.startswith("## "):
                clean_heading = self._format_markdown_for_reportlab(_clean_math_for_text(line_str[3:]))
                story.append(Paragraph(clean_heading, h1_style))
            elif line_str.startswith("# "):
                clean_heading = self._format_markdown_for_reportlab(_clean_math_for_text(line_str[2:]))
                story.append(Paragraph(clean_heading, title_style))
            elif line_str.startswith("- ") or line_str.startswith("* ") or line_str.startswith("• "):
                bullet_raw = _clean_math_for_text(line_str[2:].strip())
                bullet_xml = f"• {self._format_markdown_for_reportlab(bullet_raw)}"
                story.append(Paragraph(bullet_xml, bullet_style))
            else:
                cleaned_line = _clean_math_for_text(line_str)
                formatted_xml = self._format_markdown_for_reportlab(cleaned_line)
                story.append(Paragraph(formatted_xml, body_style))

        flush_table()
        doc.build(story, canvasmaker=NumberedCanvas)
        pdf_bytes = buffer.getvalue()
        buffer.close()

        with open(file_path, "wb") as f:
            f.write(pdf_bytes)

    @staticmethod
    def _format_markdown_for_reportlab(text: str) -> str:
        """Converts raw markdown formatting (**bold**, *italic*, `code`) into ReportLab XML."""
        escaped = (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
        escaped = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<i>\1</i>", escaped)
        escaped = re.sub(r"`(.+?)`", r'<font face="Courier" color="#8C6514">\1</font>', escaped)
        return escaped

    def generate_docx(
        self,
        title: str,
        content: str,
        author: str = "Sastra AI Assistant",
    ) -> Dict[str, Any]:
        """Generates a styled Microsoft Word (.docx) document with Khmer complex script font & clean math typography."""
        doc_id = f"doc_{uuid.uuid4().hex[:10]}"
        resolved_title = self._clean_or_extract_title(title, content)
        filename = f"{self._sanitize_filename(resolved_title)}.docx"
        file_path = GENERATED_DOCS_DIR / f"{doc_id}.docx"

        doc = docx.Document()

        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        def set_run_font(run, font_name="Noto Sans Khmer"):
            run.font.name = font_name
            rPr = run._r.get_or_add_rPr()
            rFonts = parse_xml(r"""<w:rFonts {} w:ascii="{}" w:hAnsi="{}" w:cs="{}"/>""".format(
                nsdecls("w"), font_name, font_name, font_name
            ))
            rPr.append(rFonts)

        # Document Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run(resolved_title)
        set_run_font(run_title, "Noto Sans Khmer")
        run_title.font.size = Pt(18)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(146, 64, 14)  # Gold/Amber

        # Subtitle
        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(f"Authored by {author}  •  {time.strftime('%B %d, %Y')}")
        set_run_font(run_sub, "Noto Sans Khmer")
        run_sub.font.size = Pt(9.5)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(120, 113, 108)

        # Golden separator line
        p_hr = doc.add_paragraph()
        pPr = p_hr._p.get_or_add_pPr()
        pBdr = parse_xml(r"""<w:pBdr {}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="B45309"/></w:pBdr>""".format(nsdecls("w")))
        pPr.append(pBdr)

        lines = content.strip().split("\n")
        table_buffer: List[str] = []

        def flush_docx_table():
            if not table_buffer:
                return
            rows_data = []
            for t_line in table_buffer:
                if re.match(r"^\|?\s*[-:]+\s*\|", t_line):
                    continue
                cells = [c.strip() for c in t_line.split("|")[1:-1]]
                if cells:
                    rows_data.append(cells)

            if rows_data:
                col_count = len(rows_data[0])
                table = doc.add_table(rows=len(rows_data), cols=col_count)
                table.autofit = True
                for r_idx, row in enumerate(rows_data):
                    for c_idx, val in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cleaned_val = _clean_math_for_text(self._strip_markdown(val))
                        cell.text = cleaned_val
                        if r_idx == 0:
                            shading_elm = parse_xml(r'<w:shd {} w:fill="92400E"/>'.format(nsdecls('w')))
                            cell._tc.get_or_add_tcPr().append(shading_elm)
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                                    set_run_font(run, "Noto Sans Khmer")
                        else:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    set_run_font(run, "Noto Sans Khmer")
                doc.add_paragraph()
            table_buffer.clear()

        for line in lines:
            line_str = line.strip()
            if line_str.startswith("|") and line_str.endswith("|"):
                table_buffer.append(line_str)
                continue
            else:
                flush_docx_table()

            if not line_str:
                doc.add_paragraph()
            elif re.match(r"^[-*_]{3,}$", line_str):
                # Divider border
                p_sep = doc.add_paragraph()
                pPr_sep = p_sep._p.get_or_add_pPr()
                pBdr_sep = parse_xml(r"""<w:pBdr {}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E7E5E4"/></w:pBdr>""".format(nsdecls("w")))
                pPr_sep.append(pBdr_sep)
            elif line_str.startswith("### "):
                p = doc.add_paragraph()
                r = p.add_run(_clean_math_for_text(self._strip_markdown(line_str[4:])))
                set_run_font(r, "Noto Sans Khmer")
                r.font.size = Pt(11.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(28, 25, 23)
            elif line_str.startswith("## "):
                p = doc.add_paragraph()
                r = p.add_run(_clean_math_for_text(self._strip_markdown(line_str[3:])))
                set_run_font(r, "Noto Sans Khmer")
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(180, 83, 9)
            elif line_str.startswith("# "):
                p = doc.add_paragraph()
                r = p.add_run(_clean_math_for_text(self._strip_markdown(line_str[2:])))
                set_run_font(r, "Noto Sans Khmer")
                r.font.size = Pt(15)
                r.font.bold = True
                r.font.color.rgb = RGBColor(146, 64, 14)
            elif line_str.startswith("- ") or line_str.startswith("* ") or line_str.startswith("• "):
                p = doc.add_paragraph(style="List Bullet")
                self._add_formatted_runs_docx(p, _clean_math_for_text(line_str[2:]), set_run_font)
            else:
                cleaned = _clean_math_for_text(line_str)
                p = doc.add_paragraph()
                self._add_formatted_runs_docx(p, cleaned, set_run_font)

        flush_docx_table()
        doc.save(str(file_path))
        file_size = file_path.stat().st_size

        meta = {
            "id": doc_id,
            "title": resolved_title,
            "format": "docx",
            "filename": filename,
            "size_bytes": file_size,
            "size_formatted": self._format_bytes(file_size),
            "created_at": time.time(),
            "download_url": f"/api/documents/generated/{doc_id}",
        }
        self._meta[doc_id] = meta
        self._save_meta()
        return meta

    @staticmethod
    def _strip_markdown(text: str) -> str:
        return re.sub(r"[*_`#]", "", text).strip()

    @staticmethod
    def _add_formatted_runs_docx(paragraph, text: str, set_font_fn):
        """Parse **bold** markdown into docx runs with Khmer font and styling applied."""
        parts = re.split(r"(\*\*.*?\*\*)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            else:
                run = paragraph.add_run(part)
            set_font_fn(run, "Noto Sans Khmer")

    def generate_markdown(self, title: str, content: str) -> Dict[str, Any]:
        """Generates a clean Markdown file with YAML frontmatter."""
        doc_id = f"doc_{uuid.uuid4().hex[:10]}"
        resolved_title = self._clean_or_extract_title(title, content)
        filename = f"{self._sanitize_filename(resolved_title)}.md"
        file_path = GENERATED_DOCS_DIR / f"{doc_id}.md"

        full_md = f"""---
title: "{resolved_title}"
author: "Sastra AI"
date: "{time.strftime('%Y-%m-%d %H:%M:%S UTC')}"
generator: "Sastra AI Sovereign Assistant"
---

# {resolved_title}

{content.strip()}
"""
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(full_md)

        file_size = file_path.stat().st_size
        meta = {
            "id": doc_id,
            "title": resolved_title,
            "format": "md",
            "filename": filename,
            "size_bytes": file_size,
            "size_formatted": self._format_bytes(file_size),
            "created_at": time.time(),
            "download_url": f"/api/documents/generated/{doc_id}",
        }
        self._meta[doc_id] = meta
        self._save_meta()
        return meta

    def generate_csv(self, title: str, content: str) -> Dict[str, Any]:
        """Generates a structured CSV file."""
        doc_id = f"doc_{uuid.uuid4().hex[:10]}"
        resolved_title = self._clean_or_extract_title(title, content)
        filename = f"{self._sanitize_filename(resolved_title)}.csv"
        file_path = GENERATED_DOCS_DIR / f"{doc_id}.csv"

        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content.strip())

        file_size = file_path.stat().st_size
        meta = {
            "id": doc_id,
            "title": resolved_title,
            "format": "csv",
            "filename": filename,
            "size_bytes": file_size,
            "size_formatted": self._format_bytes(file_size),
            "created_at": time.time(),
            "download_url": f"/api/documents/generated/{doc_id}",
        }
        self._meta[doc_id] = meta
        self._save_meta()
        return meta

    def get_document(self, doc_id: str) -> Optional[Dict[str, Any]]:
        return self._meta.get(doc_id)

    def get_document_path(self, doc_id: str) -> Optional[Path]:
        meta = self._meta.get(doc_id)
        if not meta:
            return None
        fmt = meta.get("format", "pdf")
        path = GENERATED_DOCS_DIR / f"{doc_id}.{fmt}"
        return path if path.exists() else None

    @staticmethod
    def _sanitize_filename(name: str) -> str:
        s = re.sub(r"[^\w\s\-\(\)]", "", name, flags=re.UNICODE).strip()
        s = re.sub(r"\s+", "_", s)
        return s[:60] or "sastra_document"

    @staticmethod
    def _format_bytes(size: int) -> str:
        if size < 1024:
            return f"{size} B"
        if size < 1024 * 1024:
            return f"{size / 1024:.1f} KB"
        return f"{size / (1024 * 1024):.1f} MB"


doc_generator = DocumentGenerator()
